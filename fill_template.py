#!/usr/bin/env python3
"""
fill_template.py — populate a website-copy CCT (Content Creation Template)
Word document from a "generated_content.json" package.

USAGE:
    python3 fill_template.py <template.docx> <content.json> <output.docx>
"""
import copy
import html
import itertools
import json
import re
import sys
from pathlib import Path

import docx
from docx.document import Document as DocumentObj
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

W_P = qn('w:p')
W_TBL = qn('w:tbl')
W_TR = qn('w:tr')
W_TC = qn('w:tc')
W_T = qn('w:t')
W_R = qn('w:r')

INTEXT_BUTTON_POOL = [
    "Contact Us",
    "Contact Us Now",
    "Contact Us Today",
    "Enquire Now",
    "Contact Our Team",
    "Contact Our Experts",
    "Enquire Today",
]

LEAD_BUTTON_POOL = [
    "Get In Touch",
    "Enquire Now",
    "Get A Quote",
    "Have Any Questions?",
    "Find Out More",
    "Enquire Today",
    "Reach Your Local Expert",
    "Start Your Enquiry",
    "Ask A Question",
]


def make_picker(pool):
    cycle = itertools.cycle(pool)

    def pick(avoid=None):
        val = next(cycle)
        if avoid is not None and val == avoid:
            val = next(cycle)
        return val

    return pick


def scratch_parent(elements):
    parent = OxmlElement('w:body')
    for e in elements:
        parent.append(e)
    return parent


def release(parent):
    return list(parent)


def clone(elm):
    return copy.deepcopy(elm)


def paragraph_text(p_elm):
    return "".join(t.text or "" for t in p_elm.iter(W_T))


BOLD_MARKUP_RE = re.compile(r'\*\*(.+?)\*\*')

_RPR_ORDER = ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps',
              'strike', 'dstrike', 'outline', 'shadow', 'emboss', 'imprint',
              'color', 'spacing', 'w', 'kern', 'position', 'sz', 'szCs',
              'highlight', 'u', 'effect', 'bdr', 'shd', 'vertAlign', 'rtl', 'cs',
              'lang']


def _insert_rpr_child(rPr, new_elm, tag_localname):
    idx_of_new = _RPR_ORDER.index(tag_localname)
    insert_pos = len(list(rPr))
    for i, child in enumerate(rPr):
        local = child.tag.split('}')[-1]
        if local in _RPR_ORDER and _RPR_ORDER.index(local) > idx_of_new:
            insert_pos = i
            break
    rPr.insert(insert_pos, new_elm)


def _clear_bold(rPr):
    b = rPr.find(qn('w:b'))
    if b is not None:
        rPr.remove(b)
    bcs = rPr.find(qn('w:bCs'))
    if bcs is not None:
        rPr.remove(bcs)


def _set_bold(rPr):
    b = rPr.find(qn('w:b'))
    if b is None:
        b = OxmlElement('w:b')
        _insert_rpr_child(rPr, b, 'b')
    elif qn('w:val') in b.attrib:
        del b.attrib[qn('w:val')]


def _append_markup_run(p_elm, template_run, text, bold):
    new_run = clone(template_run)
    rPr = new_run.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        new_run.insert(0, rPr)
    _set_bold(rPr) if bold else _clear_bold(rPr)
    t_elems = new_run.findall(W_T)
    if not t_elems:
        t_new = new_run.makeelement(W_T, {})
        new_run.append(t_new)
        t_elems = [t_new]
    t_elems[0].text = text
    t_elems[0].set(qn('xml:space'), 'preserve')
    for t in t_elems[1:]:
        new_run.remove(t)
    p_elm.append(new_run)
    return new_run


def set_paragraph_text(p_elm, text):
    runs = p_elm.findall(W_R)
    if not runs:
        # Create a bare run if none exists
        r = p_elm.makeelement(W_R, {})
        t = r.makeelement(W_T, {})
        t.text = text
        t.set(qn('xml:space'), 'preserve')
        r.append(t)
        p_elm.append(r)
        return
    template_run = runs[0]

    if '**' not in text:
        for r in runs[1:]:
            p_elm.remove(r)
        t_elems = template_run.findall(W_T)
        if not t_elems:
            t_new = template_run.makeelement(W_T, {})
            template_run.append(t_new)
            t_elems = [t_new]
        t_elems[0].text = text
        t_elems[0].set(qn('xml:space'), 'preserve')
        for t in t_elems[1:]:
            template_run.remove(t)
        return

    for r in runs:
        p_elm.remove(r)

    pos = 0
    any_span = False
    for m in BOLD_MARKUP_RE.finditer(text):
        any_span = True
        if m.start() > pos:
            _append_markup_run(p_elm, template_run, text[pos:m.start()], bold=False)
        _append_markup_run(p_elm, template_run, m.group(1), bold=True)
        pos = m.end()
    if pos < len(text) or not any_span:
        _append_markup_run(p_elm, template_run, text[pos:], bold=False)


def set_paragraph_text_multiline(p_elm, text):
    lines = [ln.strip() for ln in text.split("\n") if ln.strip() != ""]
    if not lines:
        set_paragraph_text(p_elm, "")
        return [p_elm]
    set_paragraph_text(p_elm, lines[0])
    produced = [p_elm]
    anchor = p_elm
    for line in lines[1:]:
        new_p = clone(p_elm)
        set_paragraph_text(new_p, line)
        anchor.addnext(new_p)
        anchor = new_p
        produced.append(new_p)
    return produced


def replace_run_substring(p_elm, exact_old, new_text):
    for r in p_elm.findall(W_R):
        t_elems = r.findall(W_T)
        if len(t_elems) == 1 and (t_elems[0].text or "") == exact_old:
            t_elems[0].text = new_text
            t_elems[0].set(qn('xml:space'), 'preserve')
            return True
    return False


def set_cell_text(cell: _Cell, text):
    p_elms = cell._tc.findall(W_P)
    if not p_elms:
        return
    first_p = p_elms[0]
    for p in p_elms[1:]:
        cell._tc.remove(p)
    if not first_p.findall(W_R):
        r = first_p.makeelement(W_R, {})
        t = r.makeelement(W_T, {})
        t.text = ""
        r.append(t)
        first_p.append(r)
    set_paragraph_text_multiline(first_p, text)


def find_tables(body):
    return body.findall(W_TBL)


def is_page_header_table(tbl_elm):
    trs = tbl_elm.findall(W_TR)
    if len(trs) != 4:
        return False
    row1_cells = trs[1].findall(W_TC)
    row3_cells = trs[3].findall(W_TC)
    if len(row1_cells) < 1 or len(row3_cells) < 1:
        return False
    r1c0 = "".join(t.text or "" for t in row1_cells[0].iter(W_T)).strip()
    r3c0 = "".join(t.text or "" for t in row3_cells[0].iter(W_T)).strip()
    return r1c0 in ("Page Title", "Page/Meta Title") and r3c0.startswith("H1")


def table_first_cell_text(tbl_elm):
    trs = tbl_elm.findall(W_TR)
    tcs = trs[0].findall(W_TC)
    return "".join(t.text or "" for t in tcs[0].iter(W_T)).strip()


def classify_block_label(label):
    low = label.lower()
    if low.startswith("home"):
        return "home"
    if low.startswith("contact"):
        return "contact"
    if low.startswith("gallery"):
        return "gallery"
    if low.startswith("services") or low.startswith("service"):
        return "generic"
    return "generic"


def extract_blocks(template_path):
    doc = docx.Document(template_path)
    body = doc.element.body
    children = list(body)
    header_positions = [i for i, ch in enumerate(children)
                         if ch.tag == W_TBL and is_page_header_table(ch)]
    if not header_positions:
        raise ValueError("No page-header tables found in template — cannot locate page blocks.")

    front_matter = [clone(e) for e in children[:header_positions[0]]]

    stencils = {}
    for idx, pos in enumerate(header_positions):
        end = header_positions[idx + 1] if idx + 1 < len(header_positions) else len(children)
        label = table_first_cell_text(children[pos])
        block_type = classify_block_label(label)
        if block_type not in stencils:
            stencils[block_type] = [clone(e) for e in children[pos:end]]
    return doc, front_matter, stencils


def fill_header_table(tbl_elm, doc, page_label, meta_heading, meta_description, headline):
    t = Table(tbl_elm, doc)
    set_cell_text(t.rows[0].cells[0], page_label)
    set_cell_text(t.rows[1].cells[1], meta_heading or "")
    set_cell_text(t.rows[2].cells[1], meta_description or "")
    set_cell_text(t.rows[3].cells[1], headline or "")


def _force_font(rPr, ascii_font="Calibri", half_points=24):
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        _insert_rpr_child(rPr, rFonts, 'rFonts')
    for attr in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        rFonts.set(qn(f'w:{attr}'), ascii_font)
    for tag in ('sz', 'szCs'):
        el = rPr.find(qn(f'w:{tag}'))
        if el is None:
            el = OxmlElement(f'w:{tag}')
            _insert_rpr_child(rPr, el, tag)
        el.set(qn('w:val'), str(half_points))


def fill_usps_table(tbl_elm, doc, values):
    t = Table(tbl_elm, doc)
    header_tr = t.rows[0]._tr
    new_tr = clone(header_tr)
    new_row = docx.table._Row(new_tr, t)
    cells = new_row.cells
    src_runs = cells[0]._tc.findall('.//' + W_R)
    template_run = clone(src_runs[0]) if src_runs else None
    for cell, value in zip(cells, values):
        if not cell._tc.findall('.//' + W_R) and template_run is not None:
            p_elms = cell._tc.findall(W_P)
            if p_elms:
                p_elms[0].append(clone(template_run))
        for r in cell._tc.findall('.//' + W_R):
            rPr = r.find(qn('w:rPr'))
            if rPr is None:
                rPr = OxmlElement('w:rPr')
                r.insert(0, rPr)
            _force_font(rPr)
        set_cell_text(cell, value or "")
    header_tr.addnext(new_tr)


def replace_run_fuzzy(p_elm, needle, new_core_text):
    for r in p_elm.findall(W_R):
        t_elems = r.findall(W_T)
        if len(t_elems) != 1:
            continue
        raw = t_elems[0].text or ""
        if raw.strip() == needle:
            lead = raw[:len(raw) - len(raw.lstrip())]
            trail = raw[len(raw.rstrip()):]
            t_elems[0].text = f"{lead}{new_core_text}{trail}"
            t_elems[0].set(qn('xml:space'), 'preserve')
            return True
    return False


def set_run_at_index(p_elm, index, text):
    runs = p_elm.findall(W_R)
    if index >= len(runs):
        return False
    run = runs[index]
    t_elems = run.findall(W_T)
    if not t_elems:
        t_new = run.makeelement(W_T, {})
        run.append(t_new)
        t_elems = [t_new]
    t_elems[0].text = text
    t_elems[0].set(qn('xml:space'), 'preserve')
    for t in t_elems[1:]:
        run.remove(t)
    return True


def fill_cta_table(tbl_elm, doc, heading, body_text, button_text):
    t = Table(tbl_elm, doc)
    content_cell = t.rows[1].cells[0]
    paras = content_cell.paragraphs
    if len(paras) >= 3:
        set_run_at_index(paras[0]._p, 1, heading or "")
        set_run_at_index(paras[1]._p, 1, body_text or "")
        set_run_at_index(paras[2]._p, 1, button_text or "")
    else:
        # Fallback if content cell has fewer paragraphs
        set_cell_text(content_cell, f"Title: {heading or ''}\nText: {body_text or ''}\nButton: {button_text or ''}")


def fill_hero_button(p_elm, button_text):
    if not replace_run_fuzzy(p_elm, "Capital Case", (button_text or "").strip()):
        # Try finding "Capital Case" within run text
        for r in p_elm.findall(W_R):
            for t in r.findall(W_T):
                if "Capital Case" in (t.text or ""):
                    t.text = t.text.replace("Capital Case", (button_text or "").strip())
                    t.set(qn('xml:space'), 'preserve')


def fill_remaining_bracket_buttons(elements, text_source):
    filled = 0
    for e in elements:
        if e.tag == W_P and "[] > Button" in paragraph_text(e):
            value = text_source() if callable(text_source) else text_source
            if replace_run_fuzzy(e, "[]", f"[{value}]") or replace_run_fuzzy(e, "[] > Button", f"[{value}] > Button"):
                filled += 1
            else:
                for r in e.findall(W_R):
                    for t in r.findall(W_T):
                        if "[]" in (t.text or ""):
                            t.text = t.text.replace("[]", f"[{value}]")
                            filled += 1
    return filled


def find_para_by_exact_text(elements, text, occurrence=0):
    seen = 0
    for e in elements:
        if e.tag == W_P and paragraph_text(e).strip() == text:
            if seen == occurrence:
                return e
            seen += 1
    return None


def insert_before(anchor_elm, new_elements):
    for e in new_elements:
        anchor_elm.addprevious(e)


def remove_elm(elm):
    elm.getparent().remove(elm)


def append_heading_body(anchor_elm, unit_template, heading_text, body_text):
    unit = [clone(e) for e in unit_template]
    insert_before(anchor_elm, unit)
    set_paragraph_text(unit[0], heading_text or "")
    set_paragraph_text_multiline(unit[2] if len(unit) > 2 else unit[0], body_text or "")
    return unit


def append_heading_only(anchor_elm, heading_template_p, heading_text):
    p = clone(heading_template_p)
    insert_before(anchor_elm, [p])
    set_paragraph_text(p, heading_text or "")
    return p


def append_bare_label(anchor_elm, label_template_p, label_text):
    p = clone(label_template_p)
    insert_before(anchor_elm, [p])
    set_paragraph_text(p, label_text or "")
    return p


def append_body_only(anchor_elm, body_template_p, spacer_template_p, body_text):
    body_p = clone(body_template_p)
    spacer_p = clone(spacer_template_p)
    insert_before(anchor_elm, [body_p, spacer_p])
    set_paragraph_text_multiline(body_p, body_text or "")
    return body_p


def get(d, *keys, default=""):
    cur = d
    for k in keys:
        if not isinstance(cur, dict) or k not in cur:
            return default
        cur = cur[k]
    if cur is None:
        return default
    if isinstance(cur, str):
        return html.unescape(cur)
    return cur


def page_name(page, default=""):
    return html.unescape(page.get("PageName", default) or default)


def fill_generic_page(stencil, doc, page, page_label, pick_lead, pick_intext, usps_table_template=None):
    content = page.get("AIContent", {})
    els = stencil

    # Find header table (first tbl) and CTA table (last tbl)
    tables = [e for e in els if e.tag == W_TBL]
    header_tbl = tables[0] if tables else els[0]
    cta_tbl = tables[-1] if len(tables) > 1 else (tables[0] if tables else None)

    fill_header_table(
        header_tbl, doc, page_label,
        get(content, "meta_data", "heading"),
        get(content, "meta_data", "description"),
        get(content, "headline"),
    )

    # Subheadline & Hero CTA
    paras = [e for e in els if e.tag == W_P]
    for p in paras:
        txt = paragraph_text(p).strip()
        if "Implement SQOT" in txt or "Strapline" in txt:
            set_paragraph_text(p, get(content, "subheadline"))
        elif "Hero image button" in txt:
            hero_text = get(content, "hero_cta") or pick_lead()
            fill_hero_button(p, hero_text)

    hero_text = get(content, "hero_cta") or pick_lead()

    # Fill sections 1-4
    h2_paras = [p for p in paras if paragraph_text(p).strip() == "H2"]
    for i in range(1, 5):
        sec = get(content, f"section{i}")
        if isinstance(sec, dict) and i <= len(h2_paras):
            h_elm = h2_paras[i-1]
            set_paragraph_text(h_elm, sec.get("heading", ""))
            # Next paragraph is text
            next_idx = els.index(h_elm) + 1
            if next_idx < len(els) and els[next_idx].tag == W_P:
                set_paragraph_text_multiline(els[next_idx], sec.get("description", ""))

    fill_remaining_bracket_buttons(els, get(content, "midpage_cta") or pick_intext)

    # Append Propositions & FAQs before CTA table
    anchor = cta_tbl if cta_tbl is not None else els[-1]
    
    # Template unit for heading/body
    h2_sample = h2_paras[0] if h2_paras else paras[0]
    body_sample = els[els.index(h2_sample) + 1] if h2_sample in els and els.index(h2_sample) + 1 < len(els) else paras[0]
    spacer_sample = clone(paras[1]) if len(paras) > 1 else clone(paras[0])
    unit_template = [h2_sample, spacer_sample, body_sample, spacer_sample]

    prop_title = get(content, "proposition_title")
    if prop_title:
        append_heading_only(anchor, h2_sample, prop_title)
    for key in ("proposition1", "proposition2", "proposition3", "proposition_1", "proposition_2", "proposition_3"):
        prop = get(content, key)
        if isinstance(prop, dict):
            append_heading_body(anchor, unit_template, prop.get("heading", ""), prop.get("description", ""))

    faqs = [get(content, f"FAQs{i}") for i in (1, 2, 3)]
    faqs = [f for f in faqs if isinstance(f, dict)]
    if faqs:
        append_bare_label(anchor, h2_sample, "Frequently Asked Questions")
        for faq in faqs:
            append_heading_body(anchor, unit_template, faq.get("question", ""), faq.get("answer", ""))

    if cta_tbl is not None:
        fill_cta_table(cta_tbl, doc, get(content, "CTA", "heading"),
                       get(content, "CTA", "action"),
                       get(content, "cta_button") or pick_lead(avoid=hero_text))
    return els


def fill_home_page(stencil, doc, page, page_label, pick_lead, pick_intext):
    content = page.get("AIContent", {})
    els = stencil

    tables = [e for e in els if e.tag == W_TBL]
    header_tbl = tables[0]
    cta_tbl = tables[-1]

    fill_header_table(
        header_tbl, doc, page_label,
        get(content, "meta_data", "heading"),
        get(content, "meta_data", "description"),
        get(content, "headline"),
    )

    paras = [e for e in els if e.tag == W_P]
    hero_text = get(content, "hero_cta") or pick_lead()

    for p in paras:
        txt = paragraph_text(p).strip()
        if "Implement SQOT" in txt or "Strapline" in txt:
            set_paragraph_text(p, get(content, "subheadline"))
        elif "Hero image button" in txt:
            fill_hero_button(p, hero_text)

    # USPs table is the 2nd table in Home
    if len(tables) >= 3:
        fill_usps_table(tables[1], doc, [get(content, "USP1"), get(content, "USP2"), get(content, "USP3")])

    # Section 1
    sec1 = get(content, "section1")
    h2_paras = [p for p in paras if paragraph_text(p).strip() == "H2"]
    if isinstance(sec1, dict) and h2_paras:
        set_paragraph_text(h2_paras[0], sec1.get("heading", ""))
        idx = els.index(h2_paras[0])
        if idx + 1 < len(els):
            set_paragraph_text_multiline(els[idx + 1], sec1.get("description", ""))

    # Service title
    for p in paras:
        if "What We Offer" in paragraph_text(p) or "What we offer" in paragraph_text(p):
            service_title = get(content, "service_title")
            if service_title:
                set_paragraph_text(p, service_title)
            break

    # Services
    services = [get(content, f"service_0{i}") for i in range(1, 7)]
    services = [s for s in services if isinstance(s, dict)]
    
    # Locate service summary paragraph and service link
    svc_anchor = None
    for p in paras:
        if "Why Choose Us?" in paragraph_text(p) or "Why choose us?" in paragraph_text(p):
            svc_anchor = p
            break
    if svc_anchor is None:
        svc_anchor = cta_tbl

    # Find the service template unit
    for p in paras:
        if "[] > Relevant Service Page" in paragraph_text(p):
            idx = els.index(p)
            text_p = els[idx - 2] if idx >= 2 else els[idx - 1]
            unit_template = [text_p, clone(els[idx - 1]), p, clone(els[idx + 1] if idx + 1 < len(els) else p)]
            
            # Fill first service
            if services:
                set_paragraph_text_multiline(text_p, services[0].get("summary", ""))
                replace_run_fuzzy(p, "[]", f"[{services[0].get('headline', '')}]")
            
            # Append remaining services
            for svc in services[1:]:
                u = [clone(e) for e in unit_template]
                insert_before(svc_anchor, u)
                set_paragraph_text_multiline(u[0], svc.get("summary", ""))
                replace_run_fuzzy(u[2], "[]", f"[{svc.get('headline', '')}]")
            break

    # Section 2 under Why Choose Us
    sec2 = get(content, "section2")
    if isinstance(sec2, dict) and svc_anchor is not None and svc_anchor.tag == W_P:
        idx = els.index(svc_anchor)
        body_p = els[idx + 1] if idx + 1 < len(els) else svc_anchor
        next_elem = svc_anchor.getnext()
        append_body_only(next_elem if next_elem is not None else svc_anchor, body_p, body_p, sec2.get("description", ""))

    # Testimonials & Append process / propositions
    testimonials_p = None
    for p in paras:
        if "Testimonials" in paragraph_text(p):
            testimonials_p = p
            break
    if testimonials_p is None:
        testimonials_p = cta_tbl

    h2_sample = h2_paras[0] if h2_paras else paras[0]
    body_sample = els[els.index(h2_sample) + 1] if h2_sample in els and els.index(h2_sample) + 1 < len(els) else paras[0]
    spacer_sample = clone(paras[1]) if len(paras) > 1 else clone(paras[0])
    h2_text_unit = [h2_sample, spacer_sample, body_sample, spacer_sample]

    sec3 = get(content, "section3")
    if isinstance(sec3, dict):
        append_heading_body(testimonials_p, h2_text_unit, sec3.get("heading", ""), sec3.get("description", ""))

    process_title = get(content, "process_title")
    if process_title:
        append_heading_only(testimonials_p, h2_sample, process_title)
    for key in ("process_1", "process_2", "process_3"):
        step = get(content, key)
        if isinstance(step, dict):
            append_heading_body(testimonials_p, h2_text_unit, step.get("heading", ""), step.get("description", ""))

    proposition_title = get(content, "proposition_title")
    if proposition_title:
        append_heading_only(testimonials_p, h2_sample, proposition_title)
    for key in ("proposition_1", "proposition_2", "proposition_3"):
        prop = get(content, key)
        if isinstance(prop, dict):
            append_heading_body(testimonials_p, h2_text_unit, prop.get("heading", ""), prop.get("description", ""))

    fill_remaining_bracket_buttons(els, get(content, "midpage_cta") or pick_intext)

    fill_cta_table(cta_tbl, doc, get(content, "CTA", "heading"),
                   get(content, "CTA", "action"),
                   get(content, "cta_button") or pick_lead(avoid=hero_text))
    return els


def fill_contact_page(stencil, doc, page, page_label, usps_table_template, h2_text_unit_template, pick_lead):
    content = page.get("AIContent", {})
    els = stencil

    tables = [e for e in els if e.tag == W_TBL]
    header_tbl = tables[0]

    fill_header_table(
        header_tbl, doc, page_label,
        get(content, "meta_data", "heading"),
        get(content, "meta_data", "description"),
        get(content, "headline"),
    )

    paras = [e for e in els if e.tag == W_P]
    for p in paras:
        txt = paragraph_text(p).strip()
        if "Text in 100 words" in txt or "Welcome text" in txt:
            set_paragraph_text_multiline(p, get(content, "website_conclusion"))
        elif "Hero image button" in txt:
            fill_hero_button(p, get(content, "hero_cta") or pick_lead())

    # Form anchor
    form_anchor = None
    for p in paras:
        if "FORM HEADING" in paragraph_text(p) or "Name*" in paragraph_text(p):
            form_anchor = p
            break
    if form_anchor is None:
        form_anchor = els[-1]

    sec1 = get(content, "section1")
    if isinstance(sec1, dict) and h2_text_unit_template:
        append_heading_body(form_anchor, h2_text_unit_template,
                             sec1.get("heading", ""), sec1.get("description", ""))

    usps = [get(content, "USP1"), get(content, "USP2"), get(content, "USP3")]
    if any(usps) and usps_table_template is not None:
        usps_tbl = clone(usps_table_template)
        insert_before(form_anchor, [usps_tbl])
        fill_usps_table(usps_tbl, doc, usps)

    return els


def fill_sitemap(front_matter, doc, page_names):
    tables_in_front = [e for e in front_matter if e.tag == W_TBL]
    if not tables_in_front:
        return
    t = Table(tables_in_front[0], doc)
    sitemap_cell = t.rows[2].cells[0]
    paras = sitemap_cell.paragraphs
    idx_label = None
    for i, p in enumerate(paras):
        if "SITE MAP" in p.text:
            idx_label = i
            break
    if idx_label is None or idx_label + 1 >= len(paras):
        # Just append page names
        p_last = paras[-1]._p
        for name in page_names:
            new_p = clone(p_last)
            set_paragraph_text(new_p, name)
            p_last.addnext(new_p)
            p_last = new_p
        return
    name_template = clone(paras[idx_label + 1]._p)
    existing_name_paras = paras[idx_label + 1:]
    trailing_empty = paras[-1]._p
    for p in existing_name_paras:
        remove_elm(p._p)
    for name in page_names:
        p = clone(name_template)
        trailing_empty.addprevious(p)
        set_paragraph_text(p, name)


def build_document(template_path, json_path, output_path):
    doc, front_matter, stencils = extract_blocks(template_path)
    with open(json_path, encoding='utf-8') as f:
        data = json.load(f)
    pages = data["pages"]

    home_pages = [p for p in pages if p.get("PageType") == "home"]
    contact_pages = [p for p in pages if p.get("PageType") == "contact"]
    other_pages = [p for p in pages if p.get("PageType") not in ("home", "contact")]

    pick_lead = make_picker(LEAD_BUTTON_POOL)
    pick_intext = make_picker(INTEXT_BUTTON_POOL)

    body = doc.element.body
    for child in list(body):
        body.remove(child)

    for e in front_matter:
        body.append(e)

    sitemap_names = ((["Home"] if home_pages else []) + [page_name(p) for p in other_pages]
                      + [page_name(p, "Contact Us") for p in contact_pages])
    fill_sitemap(front_matter, doc, sitemap_names)

    if "generic" not in stencils:
        raise ValueError("Template has no generic Service-page block to use as a stencil.")
    generic_template = stencils["generic"]

    # Extract USP table template from home stencil if available
    usps_table_template = None
    if "home" in stencils:
        home_tbls = [e for e in stencils["home"] if e.tag == W_TBL]
        if len(home_tbls) >= 2:
            usps_table_template = home_tbls[1]

    # Extract sample H2/Text unit
    gen_paras = [e for e in generic_template if e.tag == W_P]
    h2_gen = [p for p in gen_paras if paragraph_text(p).strip() == "H2"]
    h2_text_unit_external = None
    if h2_gen:
        h2_p = h2_gen[0]
        idx = generic_template.index(h2_p)
        body_p = generic_template[idx + 1] if idx + 1 < len(generic_template) else gen_paras[0]
        spacer_p = clone(gen_paras[1]) if len(gen_paras) > 1 else clone(gen_paras[0])
        h2_text_unit_external = [clone(h2_p), clone(spacer_p), clone(body_p), clone(spacer_p)]

    page_number = 1
    if home_pages and "home" in stencils:
        home_stencil = [clone(e) for e in stencils["home"]]
        scratch = scratch_parent(home_stencil)
        fill_home_page(home_stencil, doc, home_pages[0], f"Home (Page {page_number})",
                        pick_lead, pick_intext)
        for e in release(scratch):
            body.append(e)
        page_number += 1
    elif home_pages:
        print("WARNING: JSON has a 'home' page but template has no Home block; skipping.", file=sys.stderr)

    for p in other_pages:
        page_stencil = [clone(e) for e in generic_template]
        scratch = scratch_parent(page_stencil)
        label = f'{page_name(p)} (Page {page_number})'
        fill_generic_page(page_stencil, doc, p, label, pick_lead, pick_intext, usps_table_template)
        for e in release(scratch):
            body.append(e)
        page_number += 1

    if contact_pages:
        if "contact" in stencils:
            contact_stencil = [clone(e) for e in stencils["contact"]]
            scratch = scratch_parent(contact_stencil)
            label = f"Contact Us (Page {page_number})"
            fill_contact_page(contact_stencil, doc, contact_pages[0], label,
                               usps_table_template, h2_text_unit_external, pick_lead)
            for e in release(scratch):
                body.append(e)
            page_number += 1
        else:
            print("WARNING: JSON has a 'contact' page but template has no Contact block; skipping.", file=sys.stderr)

    doc.save(output_path)
    return output_path


def main():
    if len(sys.argv) != 4:
        print("Usage: python3 fill_template.py <template.docx> <content.json> <output.docx>")
        sys.exit(1)
    template_path, json_path, output_path = sys.argv[1], sys.argv[2], sys.argv[3]
    build_document(template_path, json_path, output_path)
    print(f"Wrote {output_path}")


if __name__ == "__main__":
    main()
